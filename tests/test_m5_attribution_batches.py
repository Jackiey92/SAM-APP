from __future__ import annotations

import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace
from urllib.error import HTTPError
import zipfile

import pytest

from app.db import init_db, connect
from app.employees.budget import AttributionBudget, EmployeeBudgetExceeded
from app.employees.semantic import SemanticEmployee, SemanticDecisionUnavailable
from app.employees.worker import EmployeeRunner
from app.employees.model_tool import EmployeeModelTool
from app.entities.bridge import EntityBridgeService, _text_blocks
from app.harness.staging import StagingStore
from app.runtime_errors import EmployeeTransportError
from app.storage import ArchiveFileStore
from app.storage.mapping_store import render_mapping


def setup_file(tmp_path, n=53):
    db = tmp_path / 'app.db'
    init_db(db)
    stored = ArchiveFileStore(tmp_path / 'objects', db).put_bytes(b'acc-batch', original_name='acc-batch.docx')
    StagingStore(db).save_manifest({'file_hash': stored.file_hash, 'filename': 'acc-batch.docx',
        'format': 'docx', 'parse_summary': {'status': 'parsed'},
        'pages': [{'page_no': None, 'text_items': [
            {'text': f'block {i}', 'source_loc': {'locator': f'paragraph/{i}'}} for i in range(n)], 'tables': []}]})
    return db, stored


def batch_runner(calls):
    def run(**kw):
        calls.append(kw)
        blocks = json.loads(kw['text'].split('\n', 1)[1])['blocks']
        return {'candidates': [{'block_id': block['block_id'], 'label': 'ambiguous',
                               'reason': 'not enough evidence', 'needs_review': True} for block in reversed(blocks)]}
    return run


def test_53_blocks_use_two_stateless_model_batches(tmp_path):
    db, stored = setup_file(tmp_path)
    calls = []
    employee = SemanticEmployee(EmployeeRunner(runner=batch_runner(calls), retry_backoff=0))
    result = EntityBridgeService(db, employee=employee).run(company_id='acc-batch', file_hash=stored.file_hash,
                                                          thread_id='parse-job:must-not-grow')
    assert result['model_calls'] == result['batch_count'] == len(calls) == 2
    assert result['batch_sizes'] == [32, 21]
    assert all(call['scope']['thread_id'] is None for call in calls)
    rows = EntityBridgeService(db).list(company_id='acc-batch', run_id=result['run_id'])
    assert len(rows) == 53 and all(row['classification'] == 'ambiguous' for row in rows)
    assert [row['content'] for row in rows] == [f'block {i}' for i in range(53)]


def test_budget_circuit_retains_completed_prefix_and_metrics(tmp_path):
    db, stored = setup_file(tmp_path, 5)
    calls = []
    employee = SemanticEmployee(EmployeeRunner(runner=batch_runner(calls), retry_backoff=0))
    service = EntityBridgeService(db, employee=employee, batch_size=2, budget=AttributionBudget(max_calls=1))
    with pytest.raises(EmployeeBudgetExceeded, match=r'completed_blocks=2/5 model_calls=1'):
        service.run(company_id='acc-batch', file_hash=stored.file_hash)
    with connect(db) as conn:
        artifact = dict(conn.execute('SELECT * FROM entity_bridge_artifacts').fetchone())
        assert artifact['status'] == 'failed' and artifact['block_count'] == 2
        assert conn.execute('SELECT count(*) FROM entity_bridge_blocks').fetchone()[0] == 2
    assert len(calls) == 1
    metrics = service._manifest(stored.file_hash)['attribution_metrics']
    assert metrics['model_calls'] == 1 and metrics['completed_blocks'] == 2


@pytest.mark.parametrize('bad_ids', [[0], [0, 0], [0, 99], [0, True]])
def test_batch_ids_are_exact_not_positional(bad_ids):
    employee = SemanticEmployee(EmployeeRunner(runner=lambda **kw: {'candidates': [
        {'block_id': i, 'label': 'self'} for i in bad_ids]}))
    with pytest.raises(SemanticDecisionUnavailable):
        employee.attribute_blocks(blocks=[{'block_id': 0, 'text': 'x'}, {'block_id': 1, 'text': 'y'}],
                                  company_name=None, aliases=[], company_id='acc-batch')


def test_character_bound_and_oversize_keep_prior_results(tmp_path):
    db, stored = setup_file(tmp_path, 3)
    calls = []
    employee = SemanticEmployee(EmployeeRunner(runner=batch_runner(calls)))
    result = EntityBridgeService(db, employee=employee, batch_chars=14).run(company_id='acc-batch', file_hash=stored.file_hash)
    assert result['batch_sizes'] == [2, 1]
    staging = StagingStore(db)
    with connect(db) as conn:
        sid = conn.execute('SELECT id FROM parse_staging').fetchone()[0]
    payload = EntityBridgeService(db)._manifest(stored.file_hash)
    payload['pages'][0]['text_items'][1]['text'] = 'x' * 40
    staging.update_payload(sid, payload)
    with pytest.raises(EmployeeBudgetExceeded, match='completed_blocks=1/3'):
        EntityBridgeService(db, employee=employee, batch_chars=14).run(company_id='acc-batch', file_hash=stored.file_hash)


@pytest.mark.parametrize('error', [
    "Agent couldn't generate a response. Please try again.",
    'HTTP 429 Too many requests', 'HTTP 503 service unavailable', 'connection closed',
])
def test_transient_errors_retry_and_exhaustion_is_transparent(error):
    calls = []
    def run(**kw):
        calls.append(kw)
        return {'status': 'error', 'error': error, 'result': {'payloads': [], 'stopReason': 'error'}}
    with pytest.raises(EmployeeTransportError) as caught:
        EmployeeRunner(runner=run, retry_backoff=0).run(skill='document-classifier', text='x', company_id='acc-batch')
    assert len(calls) == 3
    assert 'attempts=3' in str(caught.value) and error in str(caught.value) and 'stopReason' in str(caught.value)


@pytest.mark.parametrize('error', ['HTTP 401 unauthorized', 'HTTP 400 invalid parameter schema; Please try again.',
                                   'HTTP 422 schema validation failed'])
def test_deterministic_rejections_never_retry_even_empty(error):
    calls = []
    def run(**kw):
        calls.append(kw)
        return {'status': 'error', 'error': error, 'result': {'payloads': []}}
    with pytest.raises(EmployeeTransportError, match='attempts=1'):
        EmployeeRunner(runner=run, retry_backoff=0).run(skill='document-classifier', text='x', company_id='acc-batch')
    assert len(calls) == 1


def test_bounded_tool_fallback_consumes_budget_only_after_retryable_exhaustion(monkeypatch):
    monkeypatch.setenv('SAM_EMPLOYEE_FALLBACK_MODEL', 'token-plan/backup')
    calls = []
    def execute(self, **kw):
        calls.append(kw['model'])
        if kw['model'] is None:
            raise EmployeeTransportError("Agent couldn't generate a response")
        return {'candidates': [{'block_id': 0, 'label': 'ambiguous'}]}
    monkeypatch.setattr(EmployeeModelTool, 'execute', execute)
    budget = AttributionBudget()
    result = EmployeeRunner(retry_backoff=0).run(skill='document-classifier', text='x', company_id='acc-batch',
                                                bounded_model=True, budget=budget)
    assert result[0]['label'] == 'ambiguous'
    assert calls == [None, None, None, 'token-plan/backup'] and budget.model_calls == 4


def test_model_tool_one_http_request_no_history_no_tools(monkeypatch):
    config = SimpleNamespace(model_api_key_env='ACC_KEY', model_name_env='ACC_MODEL', model_default='model',
                             model_base_url_env='ACC_URL', model_base_url='https://example.invalid/v1')
    monkeypatch.setenv('ACC_KEY', 'private')
    requests = []
    class Response:
        def __enter__(self): return self
        def __exit__(self, *_): pass
        def read(self): return json.dumps({'choices': [{'finish_reason': 'stop', 'message': {'content': '{"candidates": []}'}}]}).encode()
    def urlopen(request, timeout):
        requests.append(json.loads(request.data))
        assert timeout == 17
        return Response()
    monkeypatch.setattr('urllib.request.urlopen', urlopen)
    assert EmployeeModelTool(config)._request(skill_text='skill', text='bounded input', timeout=17) == '{"candidates": []}'
    assert len(requests) == 1 and len(requests[0]['messages']) == 2 and 'tools' not in requests[0]


def test_docx_cell_fragments_become_rows_with_headers_and_real_coordinates(tmp_path):
    import importlib.machinery
    path = Path('skills/document-ingest/scripts/bridge')
    spec = importlib.util.spec_from_loader('m5_ingest', importlib.machinery.SourceFileLoader('m5_ingest', str(path)))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    source = tmp_path / 'acc-table.docx'
    xml = '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
    <w:p><w:r><w:t>table subject</w:t></w:r></w:p><w:tbl>
    <w:tr><w:trPr><w:tblHeader/></w:trPr><w:tc><w:p><w:r><w:t>name</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>value</w:t></w:r></w:p></w:tc></w:tr>
    <w:tr><w:tc><w:p><w:r><w:t>A</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>100</w:t></w:r></w:p></w:tc></w:tr>
    <w:tr><w:tc><w:p><w:r><w:t>B</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>200</w:t></w:r></w:p></w:tc></w:tr>
    </w:tbl></w:body></w:document>'''
    with zipfile.ZipFile(source, 'w') as archive:
        archive.writestr('word/document.xml', xml)
    pages = module._docx_structure(source, 'f' * 64)
    assert len(pages[0]['text_items']) == 1
    manifest = {'pages': pages}
    blocks = _text_blocks(manifest)
    assert len(blocks) == 4
    assert blocks[2][1] == 'name / value => A / 100'
    assert 'table/row/1' in blocks[2][3] and 'table/row/2' in blocks[3][3]
    assert 'table/row/1' in render_mapping(manifest)
    assert pages[0]['page_no'] is None


def test_task_budget_is_shared_across_skills_and_reset(monkeypatch):
    from app.employees.budget import parse_model_budget, CURRENT_PARSE_BUDGET
    monkeypatch.setenv('SAM_PARSE_MAX_MODEL_CALLS', '2')
    calls = []
    runner = EmployeeRunner(runner=lambda **kw: calls.append(kw) or {'candidates': []})
    with parse_model_budget() as budget:
        for skill in ['document-classifier', 'finance-fact-extraction']:
            runner.run(skill=skill, text='x', company_id='acc-budget', thread_id='must-not-grow')
        with pytest.raises(EmployeeBudgetExceeded, match='model_calls=2 limit=2'):
            runner.run(skill='fact-brief', text='x', company_id='acc-budget')
        assert budget.model_calls == 2
        assert all(call['scope']['thread_id'] is None for call in calls)
    assert CURRENT_PARSE_BUDGET.get() is None
    with parse_model_budget() as next_task:
        assert next_task.model_calls == 0


def test_attribution_time_limit_preserves_completed_batch(tmp_path, monkeypatch):
    db, stored = setup_file(tmp_path, 5)
    now = [10.0]
    monkeypatch.setattr('app.employees.budget.time.monotonic', lambda: now[0])
    calls = []
    inner = batch_runner(calls)
    def run(**kw):
        value = inner(**kw)
        now[0] += 1
        return value
    budget = AttributionBudget(timeout=2, started=10.0)
    employee = SemanticEmployee(EmployeeRunner(runner=run, retry_backoff=0))
    with pytest.raises(EmployeeBudgetExceeded, match='completed_blocks=4/5'):
        EntityBridgeService(db, employee=employee, budget=budget, batch_size=2).run(
            company_id='acc-batch', file_hash=stored.file_hash)
    assert len(calls) == 2
    with connect(db) as conn:
        assert conn.execute('SELECT count(*) FROM entity_bridge_blocks').fetchone()[0] == 4


@pytest.mark.parametrize('status, attempts', [(401, 1), (422, 1), (429, 3), (503, 3)])
def test_structured_http_metadata_beats_generic_empty_status(status, attempts):
    calls = []
    raw = {'status': 'ok', 'result': {'payloads': [], 'httpStatus': status, 'stopReason': 'error'}}
    with pytest.raises(EmployeeTransportError) as caught:
        EmployeeRunner(runner=lambda **kw: calls.append(kw) or raw, retry_backoff=0).run(
            skill='document-classifier', text='x', company_id='acc-batch')
    assert len(calls) == attempts
    assert str(status) in str(caught.value) and f'attempts={attempts}' in str(caught.value)


def test_fallback_is_not_used_after_deterministic_401(monkeypatch):
    monkeypatch.setenv('SAM_EMPLOYEE_FALLBACK_MODEL', 'token-plan/backup')
    calls = []
    def execute(self, **kw):
        calls.append(kw)
        raise EmployeeTransportError('HTTP 401 Unauthorized', retryable=False)
    monkeypatch.setattr(EmployeeModelTool, 'execute', execute)
    with pytest.raises(EmployeeTransportError, match='attempts=1'):
        EmployeeRunner(retry_backoff=0).run(skill='document-classifier', text='x', company_id='acc-batch', bounded_model=True)
    assert len(calls) == 1 and calls[0]['model'] is None


def test_model_tool_subprocess_has_hard_deadline_and_private_stdin(monkeypatch):
    from app.runtime_config import RuntimeConfig
    import subprocess
    config = RuntimeConfig.from_env()
    calls = []
    def run(cmd, **kw):
        calls.append((cmd, kw))
        raise subprocess.TimeoutExpired(cmd, kw['timeout'])
    monkeypatch.setattr('app.employees.model_tool.subprocess.run', run)
    with pytest.raises(EmployeeTransportError, match='attempts=1.*timeout'):
        EmployeeRunner(config=config, retries=0).run(skill='document-classifier', text='PRIVATE SOURCE',
                                                    company_id='acc-batch', bounded_model=True)
    cmd, kw = calls[0]
    assert 'PRIVATE SOURCE' not in ' '.join(cmd)
    assert 'PRIVATE SOURCE' in kw['input'] and kw['timeout'] <= 240


def test_transport_runtime_error_retry_path():
    calls = []
    def run(**kw):
        calls.append(kw)
        raise RuntimeError("Agent couldn't generate a response. Please try again.")
    with pytest.raises(EmployeeTransportError, match='attempts=3'):
        EmployeeRunner(runner=run, retry_backoff=0).run(skill='document-classifier', text='x', company_id='acc-batch')
    assert len(calls) == 3


def test_docx_normalization_through_real_docling(tmp_path):
    import importlib.machinery
    from docx import Document
    document = Document()
    document.add_paragraph('acc-first')
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [('item', 'amount'), ('acc-company', '100')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph('acc-last')
    source = tmp_path / 'acc-structural.docx'
    document.save(source)
    spec = importlib.util.spec_from_loader('m5_real_ingest', importlib.machinery.SourceFileLoader(
        'm5_real_ingest', 'skills/document-ingest/scripts/bridge'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    pages, summary = module._run_docling(source, 'f' * 64, 'docx', source.name)
    assert summary['status'] == 'parsed' and summary['table_row_count'] == 2
    blocks = _text_blocks({'pages': pages})
    assert [block[1] for block in blocks] == ['acc-first', 'item / amount', 'acc-company / 100', 'acc-last']
    assert all(block[2] is None for block in blocks)


def test_http_parameter_root_and_json_credentials_are_preserved_safely(monkeypatch):
    import io
    config = SimpleNamespace(model_api_key_env='ACC_KEY', model_name_env='ACC_MODEL', model_default='model',
                             model_base_url_env='ACC_URL', model_base_url='https://example.invalid/v1')
    monkeypatch.setenv('ACC_KEY', 'private')
    def urlopen(*args, **kwargs):
        raise HTTPError('https://example.invalid/v1', 400, 'Bad Request', {},
                        io.BytesIO(b'{"error":"invalid parameter max_tokens", "api_key":"private-secret"}'))
    monkeypatch.setattr('urllib.request.urlopen', urlopen)
    with pytest.raises(EmployeeTransportError) as caught:
        EmployeeModelTool(config)._request(skill_text='skill', text='x', timeout=17)
    assert not caught.value.retryable
    assert 'HTTP 400' in str(caught.value) and 'max_tokens' in str(caught.value)
    assert 'private-secret' not in str(caught.value)


@pytest.mark.parametrize('raw', [None, '', '   '])
def test_empty_channel_response_is_transport_failure(raw):
    calls = []
    with pytest.raises(EmployeeTransportError, match='attempts=3.*empty response'):
        EmployeeRunner(runner=lambda **kw: calls.append(kw) or raw, retry_backoff=0).run(
            skill='document-classifier', text='x', company_id='acc-batch')
    assert len(calls) == 3


@pytest.mark.parametrize('failure', ['aborted', 'disconnected', 'incomplete', 'ssl_eof'])
def test_interrupted_http_stream_is_retryable(failure):
    import http.client
    import ssl
    errors = {'aborted': ConnectionAbortedError('connection aborted'),
              'disconnected': http.client.RemoteDisconnected('connection closed'),
              'incomplete': http.client.IncompleteRead(b'partial', 5),
              'ssl_eof': ssl.SSLEOFError('TLS stream closed')}
    calls = []
    def run(**kw):
        calls.append(kw)
        raise errors[failure]
    with pytest.raises(EmployeeTransportError, match='attempts=3'):
        EmployeeRunner(runner=run, retry_backoff=0).run(skill='document-classifier', text='x', company_id='acc-batch')
    assert len(calls) == 3
