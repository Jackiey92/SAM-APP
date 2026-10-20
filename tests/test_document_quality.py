"""Real bwrap and pinned-library security/quality regression, not mocked isolation."""
from dataclasses import replace
import hashlib
import importlib.util
import json
from pathlib import Path
import time

from docx import Document
import pytest

from app.document_quality.sandbox import QualitySandbox, SandboxLimits, SandboxFailure, conclusion
from app.document_quality.gate import DocumentQualityGate, QualityNotPassed
from app.runtime_config import RuntimeConfig
from app.db import init_db
from app.storage.archive_store import ArchiveFileStore
from app.harness.staging import StagingStore

ROOT = Path(__file__).resolve().parents[1]
VENV = ROOT / '.sam-isolated/venv'


@pytest.fixture
def sandbox(tmp_path):
    assert VENV.joinpath('bin/python').is_file(), 'Install pinned QC venv with build-sam-isolated.sh --qc-only --install'
    return QualitySandbox(VENV, temporary_root=tmp_path / 'tasks')


def run(sandbox, script, evidence=None):
    return sandbox.run(role='file_processor', company_id='acc-qc', task_id='task-1',
                       script=script, evidence=evidence or {'original.txt': b'original'})


def test_real_sandbox_readonly_source_system_store_and_outside_task(sandbox, tmp_path):
    other = tmp_path / 'other-company.txt'
    other.write_text('private')
    script = f'''
import json, os
from pathlib import Path
for path in ['/evidence/original.txt', '/usr/qc-probe', '/system-probe',
             '/home/jakki/projects/startup-ai-manager/data/app.db', '/tmp/sam-m5/main.db', {str(other)!r}]:
    try:
        Path(path).write_text('corrupt')
    except OSError:
        pass
    else:
        raise AssertionError('write allowed: ' + path)
try:
    Path({str(other)!r}).read_bytes()
except OSError:
    pass
else:
    raise AssertionError('other task readable')
assert os.getuid() != 0
assert not os.environ.get('SAM_LEADER_MODEL_API_KEY')
assert Path('/evidence/original.txt').read_text() == 'original'
Path('/work/draft.txt').write_text('temporary')
print(json.dumps({{'status':'pass','issues':[]}}))
'''
    assert run(sandbox, script)['status'] == 'pass'
    assert other.read_text() == 'private'
    assert not list(sandbox.temporary_root.iterdir())


def test_real_sandbox_network_denied(sandbox):
    import os
    host_network = os.readlink('/proc/self/ns/net')
    script = 'import os;assert os.readlink("/proc/self/ns/net") != ' + repr(host_network) + '\n' + '''
import socket,json
s=socket.socket();s.settimeout(.2)
try:
    s.connect(('1.1.1.1',443))
except OSError:
    pass
else:
    raise AssertionError('external network available')
print(json.dumps({'status':'pass','issues':[]}))
'''
    assert run(sandbox, script)['status'] == 'pass'


@pytest.mark.parametrize('role', ['finance_analyst', 'legal_analyst', 'business_analyst', 'manager'])
def test_other_roles_cannot_execute(sandbox, role):
    with pytest.raises(PermissionError):
        sandbox.run(role=role, company_id='acc-qc', task_id='t', script='print(1)', evidence={})
    assert not sandbox.temporary_root.exists()


def test_cpu_killed_and_task_directory_cleaned(sandbox):
    limited = QualitySandbox(VENV, temporary_root=sandbox.temporary_root,
                             limits=SandboxLimits(cpu_seconds=1, wall_seconds=6))
    begin = time.monotonic()
    with pytest.raises(SandboxFailure, match='quota|process failed'):
        run(limited, 'while True: pass')
    assert time.monotonic() - begin < 5
    assert not list(sandbox.temporary_root.iterdir())


def test_wall_clock_killed_and_task_directory_cleaned(sandbox):
    limited = QualitySandbox(VENV, temporary_root=sandbox.temporary_root,
                             limits=SandboxLimits(wall_seconds=.2))
    with pytest.raises(SandboxFailure, match='wall-clock quota'):
        run(limited, 'import time;time.sleep(5)')
    assert not list(sandbox.temporary_root.iterdir())


def test_disk_is_aggregate_kernel_tmpfs_quota(sandbox):
    limited = QualitySandbox(VENV, temporary_root=sandbox.temporary_root,
                             limits=SandboxLimits(disk_bytes=1024 * 1024))
    script = '''
from pathlib import Path
for i in range(20):
    Path('/work/' + str(i)).write_bytes(b'x' * 100000)
'''
    with pytest.raises(SandboxFailure, match='No space left'):
        run(limited, script)
    assert not list(sandbox.temporary_root.iterdir())


def test_output_flood_is_killed_and_cleaned(sandbox):
    limited = QualitySandbox(VENV, temporary_root=sandbox.temporary_root,
                             limits=SandboxLimits(output_bytes=2048))
    with pytest.raises(SandboxFailure, match='output quota'):
        run(limited, 'print("x" * 10000)')
    assert not list(sandbox.temporary_root.iterdir())


def test_pinned_document_libraries_available_offline(sandbox):
    script = '''
import json
from importlib.metadata import version
import docx,openpyxl,pptx,pdfplumber,pymupdf,PIL
expected={'python-docx':'1.2.0','openpyxl':'3.1.5','python-pptx':'1.0.2',
          'pdfplumber':'0.11.9','PyMuPDF':'1.28.2','Pillow':'12.3.0'}
for name,value in expected.items(): assert version(name)==value
print(json.dumps({'status':'pass','issues':[]}))
'''
    assert run(sandbox, script)['status'] == 'pass'


@pytest.fixture
def real_document(tmp_path):
    source = tmp_path / 'acc-quality.docx'
    doc = Document()
    doc.add_heading('资料', 1)
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['项目','2023年','2024年']): cell.text = value
    table.cell(1,0).text = '指标'
    cell = table.cell(1,1)
    cell.text = ''
    cell.paragraphs[0].add_run('150')
    cell.paragraphs[0].add_run('99.24')
    table.cell(1,2).text = '11413.05'
    doc.save(source)
    spec = importlib.util.spec_from_file_location('qc_fidelity_fixture', ROOT / 'skills/document-ingest/scripts/docx_fidelity.py')
    module = importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    artifacts = module.project_docx(source, digest, tmp_path / 'objects')
    manifest = {'file_hash':digest,'filename':source.name,'format':'docx', 'pages':[],
                'parse_summary':{'status':'parsed'},'document':artifacts}
    return source, manifest


VERIFY_SCRIPT = '''
import json
from pathlib import Path
from docx import Document
source = Document('/evidence/original.docx')
expected = source.tables[0].cell(1,1).text
md = Path('/evidence/document.md').read_text()
assert expected == '15099.24'
correct = '| 指标 | ' + expected + ' | 11413.05 |'
issues=[] if correct in md else [{'source_loc':'docx:body/1/table/row/1',
    'problem':'原件第三年数值被错误交换或拆开','expected':'保持原件单元格15099.24完整且位于2023年列'}]
print(json.dumps({'status':'rework' if issues else 'pass','issues':issues},ensure_ascii=False))
'''


def scripted_employee():
    def employee(**kw):
        task = json.loads(kw['text'])
        if 'last_script_result' not in task:
            return {'candidates':[{'action':'verify','script':VERIFY_SCRIPT}]}
        return {'candidates':[{'action':'conclude',**task['last_script_result']}]}
    return employee


def test_quality_direct_original_split_runs_rework_and_pass(sandbox, real_document):
    source, manifest = real_document
    gate = DocumentQualityGate(sandbox=sandbox,employee=scripted_employee(),max_reworks=1)
    manifest['document']['markdown'] = manifest['document']['markdown'].replace('15099.24 | 11413.05', '11413.05 | 15099.24')
    reworks = []
    def reparse(issues):
        reworks.append(issues)
        manifest['document']['markdown'] = manifest['document']['markdown'].replace('11413.05 | 15099.24', '15099.24 | 11413.05')
        return manifest
    result = gate.inspect(company_id='acc-qc',task_id='positive',source=source,manifest=manifest,reparse=reparse)
    assert result['status'] == 'pass'
    assert result['history'][0]['status'] == 'rework'
    assert result['history'][0]['issues'][0]['source_loc'] == 'docx:body/1/table/row/1'
    assert result['metrics'] == {'phase':'quality','model_calls':4,'script_calls':2,'reworks':1}
    assert len(reworks) == 1 and not list(sandbox.temporary_root.iterdir())


def test_repeated_rework_is_needs_human_and_no_unlimited_loop(sandbox,real_document):
    source, manifest = real_document
    manifest['document']['markdown'] = manifest['document']['markdown'].replace('15099.24','150 99.24')
    rounds=[]
    gate=DocumentQualityGate(sandbox=sandbox,employee=scripted_employee(),max_reworks=1)
    result=gate.inspect(company_id='acc-qc',task_id='bad',source=source,manifest=manifest,
                        reparse=lambda issues: rounds.append(issues) or manifest)
    assert result['status']=='needs_human' and result['retry_exhausted']
    assert len(rounds)==1 and result['metrics']['model_calls']==4


def test_budget_and_unverified_conclusion_fail_closed(sandbox,real_document):
    source, manifest = real_document
    gate=DocumentQualityGate(sandbox=sandbox,employee=scripted_employee(),max_model_calls=1)
    result=gate.inspect(company_id='acc-qc',task_id='cap',source=source,manifest=manifest)
    assert result['status']=='needs_human' and result['metrics']['model_calls']==1
    gate=DocumentQualityGate(sandbox=sandbox,employee=lambda **kw:{'candidates':[{'action':'conclude','status':'pass','issues':[]}]})
    assert gate.inspect(company_id='acc-qc',task_id='no-read',source=source,manifest=manifest)['status']=='needs_human'


def test_conclusions_reject_business_payload_or_unlocated_rework():
    with pytest.raises(ValueError): conclusion({'status':'pass','issues':[], 'facts':[{'value':1}]})
    with pytest.raises(ValueError): conclusion({'status':'rework','issues':[{'problem':'x'}]})


def test_global_code_mode_unchanged():
    path=ROOT/'harness-openclaw/state/openclaw.json'
    if not path.is_file(): path=ROOT/'harness-openclaw/state/openclaw.example.json'
    config=json.loads(path.read_text())
    assert config['tools']['codeMode']['enabled'] is False
    from app.contracts.phase3 import build_phase3_role_matrix
    matrix=build_phase3_role_matrix(ROOT/'harness-openclaw/tool-registry/manifest.json',ROOT/'skills')
    for role in ('finance_analyst','legal_analyst','business_analyst'):
        assert matrix['tool_search']['code_mode'] is False
        assert 'sam_quality_execute' not in matrix['roles'][role]['tools']


def test_ingestion_persists_canonical_md_and_sidecars(tmp_path, real_document):
    source, manifest = real_document
    db = tmp_path/'app.db';init_db(db)
    stored=ArchiveFileStore(tmp_path/'2a/bin',db).put_bytes(source.read_bytes(),original_name='acc-qc.docx')
    staging=StagingStore(db,maps_path=tmp_path/'2a/map')
    staging.save_manifest(manifest)
    path=staging.maps_path/stored.file_hash[:2]/(stored.file_hash+'.md')
    assert path.read_text()==manifest['document']['markdown']
    for name in ('document_map','assets','self_check','hyperlinks'):
        assert json.loads(path.with_suffix('.'+name+'.json').read_text())==manifest['document'][name]


def test_quality_failure_stops_before_downstream_and_job_needs_human(tmp_path,sandbox,real_document):
    source, manifest=real_document
    manifest['document']['markdown']=manifest['document']['markdown'].replace('15099.24','150 99.24')
    db=tmp_path/'app.db';init_db(db)
    objects=tmp_path/'2a/bin'
    stored=ArchiveFileStore(objects,db).put_bytes(source.read_bytes(),original_name='acc-qc.docx')
    staging=StagingStore(db,maps_path=tmp_path/'2a/map');sid=staging.save_manifest(manifest)
    from app.document_quality.service import review_parsed_document
    from app.parse_jobs import ParseJobManager
    downstream=[]
    gate=DocumentQualityGate(sandbox=sandbox,employee=scripted_employee(),max_reworks=0)
    def worker(job,progress,event):
        progress(stage='reviewing',message='质量核验')
        review_parsed_document(job=job,adapter=None,staging=staging,staging_id=sid,
            config=RuntimeConfig.from_env(project_root=ROOT),db_path=db,objects_dir=objects,
            timeout=10,gate=gate,snapshot_renderer=lambda **kw:{'pages':[],'warnings':[]})
        downstream.append('should not execute')
    manager=ParseJobManager(db,worker,workers=1)
    try:
        job=manager.create(company_id='acc-qc',file_hash=stored.file_hash,original_name='acc-qc.docx',
                           file_format='docx',harness_format='docx',size_bytes=source.stat().st_size)
        manager.executor.shutdown(wait=True)
        terminal=manager.get(job['job_id'])
        assert terminal['status']=='needs_human' and terminal['error_kind']=='document_quality'
        assert terminal['l2_committed']==0 and downstream==[]
        assert terminal['quality_status']=='needs_human' and terminal['quality_script_calls']==1
        assert terminal['quality_model_calls']==2
        assert json.loads(terminal['quality_issues'])[0]['source_loc']=='docx:body/1/table/row/1'
        record=staging.get(sid)['payload']['document_quality']
        assert record['history'][0]['status']=='rework'
        assert record['metrics']['script_calls']==1
        path=staging.maps_path/stored.file_hash[:2]/(stored.file_hash+'.quality.json')
        assert json.loads(path.read_text())['status']=='needs_human'
    finally:
        manager.close()


def test_page_snapshot_pdf_render_is_offline_and_readonly(tmp_path,sandbox):
    import pymupdf
    from app.document_quality.snapshots import render_original_pages
    source=tmp_path/'acc-page.pdf'
    with pymupdf.open() as doc:
        page=doc.new_page();page.insert_text((72,72),'Document original snapshot')
        doc.save(source)
    digest=hashlib.sha256(source.read_bytes()).hexdigest()
    result=render_original_pages(sandbox=sandbox,source=source,file_hash=digest,company_id='acc-qc',
                                task_id='snapshot',maps_path=tmp_path/'map')
    assert result['page_count']==1 and len(result['pages'])==1 and result['warnings']==[]
    png=tmp_path/'map'/digest[:2]/(digest+'.pages')/'page-001.png'
    assert png.read_bytes().startswith(b'\x89PNG')
    assert not list(sandbox.temporary_root.iterdir())


def test_quality_prompt_has_no_developer_codenames():
    text=(ROOT/'skills/document-ingest/QUALITY.md').read_text().lower()
    for forbidden in ('l2','outbox','staging','projection','bridge'):
        assert forbidden not in text
        assert forbidden not in (ROOT/'skills/document-ingest/SKILL.md').read_text().lower()
    template=(ROOT/'webapp/templates/inbox.html').read_text()
    assert 'needs_human' in template and 'quality_model_calls' in template


def test_forked_cpu_is_counted_as_one_task(sandbox):
    limited=QualitySandbox(VENV,temporary_root=sandbox.temporary_root,
                           limits=SandboxLimits(cpu_seconds=1,wall_seconds=6))
    script='import subprocess,sys\nchildren=[subprocess.Popen([sys.executable,"-c","while True: pass"]) for _ in range(8)]\nfor child in children: child.wait()'
    with pytest.raises(SandboxFailure,match='CPU/process quota|process failed'):
        run(limited,script)
    assert not list(sandbox.temporary_root.iterdir())


def test_defect_delivery_is_needs_human_with_reproduction_evidence():
    issue={'source_loc':'docx:body/1/table','problem':'成熟库与引擎原始结构不一致且无法解释','expected':'开发复现后修复转换'}
    defect={'source_loc':issue['source_loc'],'reproduction_script':'from docx import Document\nDocument("/evidence/original.docx")',
            'raw_fragment':'原始表格片段','snapshot':'/evidence/pages/1.png'}
    result=conclusion({'status':'needs_human','issues':[issue],'defects':[defect]})
    assert result['defects'][0]['reproduction_script']
    with pytest.raises(ValueError):
        conclusion({'status':'pass','issues':[],'defects':[defect]})


def test_snapshot_transport_is_multimodal_not_host_path(monkeypatch):
    from app.employees.model_tool import EmployeeModelTool
    config=RuntimeConfig.from_env(project_root=ROOT)
    monkeypatch.setenv(config.model_api_key_env,'acc-qc-test-key')
    captured=[]
    class Response:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def read(self):return json.dumps({'choices':[{'finish_reason':'stop','message':{'content':'{"candidates":[]}'}}]}).encode()
    def urlopen(request,timeout):
        captured.append(json.loads(request.data));return Response()
    monkeypatch.setattr('urllib.request.urlopen',urlopen)
    tool=EmployeeModelTool(config)
    tool._request(skill_text='质量核验',text='本任务原件页面',timeout=3,
                  model='qwen3.8-flash',images=['data:image/png;base64,YWJj'])
    assert captured[0]['model']=='qwen3.8-flash'
    assert captured[0]['messages'][1]['content'][1]['type']=='image_url'
    with pytest.raises(ValueError):
        tool._request(skill_text='quality',text='task',timeout=3,images=['file:///home/private.png'])


def test_web_quality_refusal_precedes_semantic_employees(monkeypatch,tmp_path):
    from threading import Event
    monkeypatch.setenv('SAM_DATA_ROOT',str(tmp_path/'data'))
    monkeypatch.setenv('SAM_MAIN_DB',str(tmp_path/'main.db'))
    from app.db.migrate import main as migrate
    migrate(['--db', str(tmp_path/'main.db')])
    monkeypatch.setenv('SAM_APP_DB',str(tmp_path/'app.db'))
    import webapp.app as web
    class Adapter:
        def run_parse(self,*args,**kwargs):return 1
    class Staging:
        def get(self,*args):return {'payload':{'parse_summary':{'status':'parsed'},'pages':[]}}
    monkeypatch.setattr(web,'_PARSE_ADAPTER',Adapter())
    monkeypatch.setattr(web,'StagingStore',lambda **kwargs:Staging())
    monkeypatch.setattr(web,'_has_declared_self',lambda company:True)
    calls=[]
    def denied(*args):
        calls.append('quality')
        raise QualityNotPassed({'status':'needs_human','issues':[]})
    monkeypatch.setattr(web,'_review_parsed_document',denied)
    monkeypatch.setattr(web,'semantic_employee',lambda **kwargs:calls.append('semantic'))
    stages=[]
    job={'job_id':'t','company_id':'acc-qc','file_hash':'a'*64,'harness_format':'docx'}
    with pytest.raises(QualityNotPassed):
        web._parse_job_pipeline(job,lambda **item:stages.append(item['stage']),Event())
    assert calls==['quality'] and stages[-1]=='reviewing'


def test_g1_bwrap_exposes_only_task_evidence_not_real_stores(sandbox):
    script = '''
import json
from pathlib import Path
assert list(Path('/evidence').iterdir()) == [Path('/evidence/original.txt')]
for path in ['/tmp/sam-m5/main.db', '/tmp/sam-m5/material/fmt_C_info.docx',
             '/home/jakki/projects/startup-ai-manager/data/app.db']:
    try:
        Path(path).read_bytes()
    except OSError:
        pass
    else:
        raise AssertionError('real store visible: ' + path)
    try:
        Path(path).open('r+b')
    except OSError:
        pass
    else:
        raise AssertionError('real store writable: ' + path)
print(json.dumps({'status':'pass','issues':[]}))
'''
    assert run(sandbox, script)['status'] == 'pass'
